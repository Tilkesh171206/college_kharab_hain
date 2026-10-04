from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any


def extract_attachment_text(filename: str, content_type: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)

        if suffix == ".docx":
            from docx import Document
            doc = Document(BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)

        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None]
                    if vals:
                        parts.append(" | ".join(vals))
            return "\n".join(parts)

        if suffix in {".txt", ".csv"}:
            return data.decode("utf-8", errors="replace")

    except Exception:
        return ""

    return ""
